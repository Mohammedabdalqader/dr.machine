"""Generate the SAMPLE dataset for Mu'allim.

Run from the repo root:
    .venv/Scripts/python data/generator/generate.py

Outputs (all clearly labelled SAMPLE / fictional):
    data/sample/sample_plant_fault_log.xlsx   messy maintenance log (what a plant exports)
    data/sample/sample_plant_fault_log.csv    same log as CSV
    data/sample/machines.csv                  asset list for the sample site
    data/sample/aircore_rs_service_manual_SAMPLE.pdf
    data/sample/reference/fault_catalog_clean.json   ground truth (never imported into the app)
    eval/test_set.json                        held-out evaluation cases
"""

import csv
import json
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from catalog import ALARM_CODES, FAULTS, MANUFACTURER, MODELS
from messy_phrasing import PHRASES, TECHNICIANS
from test_cases import TEST_CASES

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "sample"
EVAL = ROOT / "eval"
SEED = 7

MACHINES = [
    dict(tag="C-01", model="RS-37", site="Sahab Plant (SAMPLE)", line="Line 1", serial="SA37-2019-0412", installed="2019-05-14"),
    dict(tag="C-02", model="RS-55", site="Sahab Plant (SAMPLE)", line="Line 2", serial="SA55-2021-0077", installed="2021-02-03"),
    dict(tag="C-03", model="RS-37", site="Sahab Plant (SAMPLE)", line="Utilities", serial="SA37-2017-0156", installed="2017-09-21"),
    dict(tag="C-04", model="RS-55", site="Sahab Plant (SAMPLE)", line="Utilities", serial="SA55-2022-0301", installed="2022-06-30"),
]
FAULT_BY_ID = {f["id"]: f for f in FAULTS}


def messy_date(d: date, rng: random.Random) -> str:
    fmt = rng.choice(["%d/%m/%Y", "%d-%m-%y", "%Y-%m-%d", "%d.%m.%Y"])
    return d.strftime(fmt)


def build_log(rng: random.Random) -> list[dict]:
    rows = []
    start = date(2024, 1, 8)
    for fault in FAULTS:
        problems, actions = PHRASES[fault["id"]]
        eligible = [m for m in MACHINES if m["model"] in fault["models"]]
        for _ in range(fault["weight"]):
            machine = rng.choice(eligible)
            d = start + timedelta(days=rng.randint(0, 600))
            tag = machine["tag"]
            # Inconsistent machine naming, as in real logs.
            tag_written = rng.choice([tag, tag.replace("-", ""), f"comp {tag[-1]}", f"{tag} ({machine['model']})"])
            hours = round(fault["minutes"] / 60 * rng.uniform(0.6, 1.8), 1)
            rows.append({
                "_date": d,
                "Date": messy_date(d, rng),
                "Machine": tag_written,
                "Problem": rng.choice(problems),
                "Action taken": rng.choice(actions) if rng.random() > 0.08 else "",
                "Parts": ", ".join(fault["parts"][:1]) if fault["parts"] and rng.random() > 0.4 else "",
                "Hrs": hours if rng.random() > 0.15 else "",
                "Tech": rng.choice(TECHNICIANS),
                "_fault": fault["id"],
                "_machine": tag,
            })
    rows.sort(key=lambda r: r["_date"])
    for i, r in enumerate(rows, start=1):
        r["Ref"] = f"WO-{1000 + i}"
    return rows


def write_log(rows: list[dict]) -> None:
    cols = ["Ref", "Date", "Machine", "Problem", "Action taken", "Parts", "Hrs", "Tech"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Breakdowns"
    ws.append(["SAMPLE DATA - fictional plant, invented for development and demos"])
    ws["A1"].font = Font(bold=True, color="C0392B")
    ws.append(cols)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in rows:
        ws.append([r[c] for c in cols])
    for col, width in zip("ABCDEFGH", [10, 12, 16, 48, 52, 30, 6, 12]):
        ws.column_dimensions[col].width = width
    wb.save(OUT / "sample_plant_fault_log.xlsx")

    with open(OUT / "sample_plant_fault_log.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_machines() -> None:
    with open(OUT / "machines.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["tag", "model", "site", "line", "serial", "installed"])
        w.writeheader()
        w.writerows(MACHINES)


def write_reference(rows: list[dict]) -> None:
    ref = OUT / "reference"
    ref.mkdir(exist_ok=True)
    catalog = []
    for f in FAULTS:
        catalog.append({k: v for k, v in f.items() if k != "weight"} | {
            "log_refs": [r["Ref"] for r in rows if r["_fault"] == f["id"]],
        })
    payload = {"_note": "SAMPLE ground truth. Used only for evaluation and seeding; never shown as real data.",
               "manufacturer": MANUFACTURER, "models": MODELS, "faults": catalog}
    (ref / "fault_catalog_clean.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def write_test_set() -> None:
    EVAL.mkdir(exist_ok=True)
    machine_model = {m["tag"]: m["model"] for m in MACHINES}
    cases = []
    for c in TEST_CASES:
        exp = FAULT_BY_ID.get(c["expected"]) if c["expected"] else None
        cases.append(c | {
            "model": machine_model[c["machine"]],
            "expected_root_cause": exp["root_cause"] if exp else None,
            "expected_component": exp["component"] if exp else None,
        })
    payload = {"_note": "SAMPLE held-out test set. expected=null means the correct answer is 'not enough data'.",
               "cases": cases}
    (EVAL / "test_set.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- manual PDF

MANUAL_PAGES = [
    # (page number, title, intro text, fault ids)
    (3, "1. General safety", None, []),
    (4, "2. Technical data", None, []),
    (5, "3. Controller and emergency stop", "The controller shows alarms as codes (see the alarm table in chapter 16). "
        "Shutdown alarms stop the compressor; warnings allow it to run.", ["F16"]),
    (6, "4. Air intake", "The intake filter protects the airend from dust. Replace the element every 2000 running hours "
        "or when the differential pressure warning appears.", ["F05"]),
    (7, "5. Oil system", "Use only SA-Synth 46. Oil and oil filter interval: 4000 running hours or 12 months.", ["F02", "F23", "F24"]),
    (8, "6. Cooling and temperature control", "Normal discharge temperature is 75-95 C. Shutdown occurs at 110 C (E101).", ["F01", "F03"]),
    (9, "7. Cooling fan and drive", "RS-37 units are belt driven. RS-55 units are direct driven through a flexible coupling.", ["F04", "F17"]),
    (10, "8. Oil separation", "Separator element replacement: 8000 running hours or when separator dP exceeds 1 bar.", ["F06", "F07"]),
    (11, "9. Inlet, minimum pressure and blowdown valves", "These valves control loading, net pressure and venting at stop.", ["F09", "F08", "F21"]),
    (12, "10. Solenoid valve and safety valve", None, ["F10", "F20"]),
    (13, "11. Sensors and controller power", "Sensors: pressure transducer 0-16 bar (4-20 mA); discharge temperature PT1000.", ["F12", "F13", "F26"]),
    (14, "12. Main motor and electrical", "Electrical work must be carried out by a qualified electrician only.", ["F14", "F15", "F22", "F18"]),
    (15, "13. Pressure settings and the air network", "Factory settings: load 8.0 bar, unload 9.0 bar, max 10 bar.", ["F28", "F11"]),
    (16, "14. Condensate and dryer", "The RS-55 includes an integrated refrigerated dryer (pressure dew point +3 C).", ["F19", "F27"]),
    (17, "15. Installation and ventilation", "Maximum ambient temperature 40 C. Cooling air flow: RS-37 4.5 m3/s, RS-55 6.0 m3/s.", ["F25"]),
    (18, "16. Alarm code table", None, []),
]


def build_manual() -> int:
    path = OUT / "aircore_rs_service_manual_SAMPLE.pdf"
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=16, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=ss["Heading3"], fontSize=11, spaceBefore=6, spaceAfter=2)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=11.5)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#C0392B"))
        canvas.drawString(20 * mm, 10 * mm, "SAMPLE MANUAL - fictional equipment, for development and demos only")
        canvas.setFillColor(colors.black)
        canvas.drawRightString(190 * mm, 10 * mm, f"AirCore RS-37 / RS-55 Service Manual  -  page {doc.page}")
        canvas.restoreState()

    story = [
        Spacer(1, 60 * mm),
        Paragraph("AirCore RS-37 / RS-55", ParagraphStyle("t", parent=h1, fontSize=28, leading=34)),
        Paragraph("Rotary screw air compressors - Service manual", ParagraphStyle("t2", parent=h1, fontSize=16)),
        Spacer(1, 10 * mm),
        Paragraph(f"Manufacturer: {MANUFACTURER}", body),
        Paragraph("<font color='#C0392B'><b>SAMPLE DOCUMENT.</b> The manufacturer, models, part numbers and procedures "
                  "are invented for software development. Do not use for real equipment.</font>", body),
        PageBreak(),
        Paragraph("Contents", h1),
    ]
    for page, title, _, _ in MANUAL_PAGES:
        story.append(Paragraph(f"{title} .......... page {page}", body))
    story.append(PageBreak())

    for page, title, intro, fault_ids in MANUAL_PAGES:
        story.append(Paragraph(title, h1))
        if page == 3:
            for line in ["Always stop the compressor and apply lockout/tagout on the main isolator before any work.",
                         "Release all stored pressure. The separator tank stays pressurized after stop until blowdown is complete.",
                         "Oil and internal parts can exceed 90 C. Allow at least 15 minutes to cool.",
                         "Electrical work only by qualified electricians. Verify zero voltage before touching terminals.",
                         "Never adjust, plug or remove the safety valve.",
                         "Wear gloves, safety glasses and hearing protection (noise up to 72 dB(A))."]:
                story.append(Paragraph(f"- {line}", body))
        elif page == 4:
            data = [["", "RS-37", "RS-55"], ["Motor power", "37 kW", "55 kW"], ["Drive", "V-belts", "Direct, coupling"],
                    ["Free air delivery at 8 bar", "6.1 m3/min", "9.4 m3/min"], ["Max working pressure", "10 bar", "10 bar"],
                    ["Safety valve setting", "11 bar", "11 bar"], ["Oil capacity", "22 L", "30 L"], ["Integrated dryer", "No", "Yes"]]
            t = Table(data, colWidths=[70 * mm, 45 * mm, 45 * mm])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 9),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD"))]))
            story.append(t)
        elif page == 18:
            data = [["Code", "Meaning", "See chapter / page"]]
            for code, meaning in ALARM_CODES.items():
                pages = sorted({FAULT_BY_ID[f]["manual_page"] for f in FAULT_BY_ID if code in FAULT_BY_ID[f]["codes"]})
                data.append([code, meaning, ", ".join(f"p. {p}" for p in pages)])
            t = Table(data, colWidths=[20 * mm, 90 * mm, 50 * mm])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 9),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD"))]))
            story.append(t)
        if intro:
            story.append(Paragraph(intro, body))
        for fid in fault_ids:
            f = FAULT_BY_ID[fid]
            assert f["manual_page"] == page, f"{fid} expected on page {f['manual_page']}, placed on {page}"
            models = "" if len(f["models"]) == 2 else f" ({', '.join(f['models'])} only)"
            codes = f" - alarm {', '.join(f['codes'])}" if f["codes"] else ""
            story.append(Paragraph(f"Troubleshooting: {f['component']}{models}{codes}", h2))
            story.append(Paragraph(f"<b>Symptom:</b> {f['symptoms']}", body))
            story.append(Paragraph(f"<b>Possible cause:</b> {f['root_cause']}", body))
            steps = "<br/>".join(f"{i}. {s}" for i, s in enumerate(f["fix_steps"], 1))
            story.append(Paragraph(f"<b>Procedure:</b><br/>{steps}", body))
            if f["parts"]:
                story.append(Paragraph(f"<b>Parts:</b> {', '.join(f['parts'])}", small))
            story.append(Paragraph(f"<b>Safety:</b> {' '.join(f['safety'])}", small))
        story.append(PageBreak())
    story.pop()  # no trailing blank page

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm, title="AirCore RS Service Manual (SAMPLE)")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return doc.page


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    rows = build_log(rng)
    write_log(rows)
    write_machines()
    write_reference(rows)
    write_test_set()
    pages = build_manual()
    expected_pages = MANUAL_PAGES[-1][0]
    if pages != expected_pages:
        raise SystemExit(f"Manual has {pages} pages, expected {expected_pages}: a section overflowed; page refs are wrong.")
    print(f"log rows: {len(rows)}, faults: {len(FAULTS)}, test cases: {len(TEST_CASES)}, manual pages: {pages}")


if __name__ == "__main__":
    main()
